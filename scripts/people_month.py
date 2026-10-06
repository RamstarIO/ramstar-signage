#!/usr/bin/env python3
"""
Make this month's birthday and work-anniversary slide entries from HR's
Active Employees spreadsheet.

    python3 scripts/people_month.py ~/Documents/ramstar-private/Active_Employees.xlsx 2026-10

Prints two slides.json entries ("birthdays" and "anniversaries") to paste
into slides.json, replacing last month's.

PRIVACY: the spreadsheet holds full birth dates (with year) and hire dates.
Keep it OUTSIDE this repo. This script only prints what the slides show:
day of the month, display name and (for anniversaries) years of service.
No birth year ever leaves the spreadsheet.

Needs openpyxl:  python3 -m pip install openpyxl
"""
import json
import sys
from datetime import date, datetime

from openpyxl import load_workbook

MAX_PER_SLIDE = 8  # both slide types fit up to 8 people

# Display fixes for names the spreadsheet stores awkwardly, keyed by
# (FirstName, Name) exactly as they appear in the file (upper case).
NAME_OVERRIDES = {
    ("JOHN", "WINKFIELD"): "John Winkfield Sr.",
    ("JOHN JR", "WINKFIELDJR"): "John Winkfield Jr.",
}

# Surname particles that stay capitalised as written below.
SMALL_PARTS = {"DA": "Da", "LE": "Le", "DE": "De", "VAN": "Van"}


def display_name(first, last):
    key = (first.strip().upper(), last.strip().upper())
    if key in NAME_OVERRIDES:
        return NAME_OVERRIDES[key]
    return f"{tidy(first)} {tidy(last)}"


def tidy(text):
    """'DA DALT' -> 'Da Dalt', 'McBRAYNE' -> 'McBrayne', 'JOHN' -> 'John'."""
    words = []
    for word in text.strip().split():
        upper = word.upper()
        if upper in SMALL_PARTS:
            words.append(SMALL_PARTS[upper])
        elif upper.startswith("MC") and len(word) > 2:
            words.append("Mc" + word[2:].capitalize())
        else:
            words.append("-".join(p.capitalize() for p in word.split("-")))
    return " ".join(words)


def parse_date(value):
    """Cells arrive as real dates or as text like '10/01/1966 00:00:00' or '03/10/1978'."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%m/%d/%Y %H:%M:%S", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def day_label(year, month, day):
    """'Thu 1' style label for this year's date (matches the birthday slide)."""
    try:
        d = date(year, month, day)
    except ValueError:  # Feb 29 in a non-leap year
        d = date(year, 2, 28)
    return f"{d.strftime('%a')} {d.day}"


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    path, ym = sys.argv[1], sys.argv[2]
    year, month = (int(p) for p in ym.split("-"))

    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]
    col = {name: header.index(name) for name in ("Name", "FirstName", "BirthDate", "HiringDate")}

    birthdays, anniversaries, warnings = [], [], []
    for row in rows:
        if not row or not row[col["Name"]]:
            continue
        name = display_name(row[col["FirstName"]], row[col["Name"]])
        born = parse_date(row[col["BirthDate"]])
        hired = parse_date(row[col["HiringDate"]])

        if born is None:
            warnings.append(f"{name}: birth date not readable")
        elif born.month == month:
            birthdays.append((born.day, name))

        if hired is None:
            warnings.append(f"{name}: hire date not readable")
        elif hired.year <= 1900:
            warnings.append(f"{name}: hire date missing (shows as {hired.year}), skipped for anniversaries")
        elif hired.month == month and year - hired.year >= 1:
            anniversaries.append((hired.day, name, year - hired.year))

    birthdays.sort()
    anniversaries.sort()
    for label, people in (("birthdays", birthdays), ("anniversaries", anniversaries)):
        if len(people) > MAX_PER_SLIDE:
            warnings.append(f"{len(people)} {label} this month; the slide fits {MAX_PER_SLIDE}")

    month_name = date(year, month, 1).strftime("%B").upper()
    entries = [
        {"title": "Birthdays", "type": "birthdays", "month": month_name,
         "people": [{"day": day_label(year, month, d), "name": n} for d, n in birthdays]},
        {"title": "Work anniversaries", "type": "anniversaries", "month": month_name,
         "people": [{"day": day_label(year, month, d), "name": n, "years": y} for d, n, y in anniversaries]},
    ]
    print(",\n".join(json.dumps(e, indent=2, ensure_ascii=False) for e in entries))
    for w in warnings:
        print("WARNING: " + w, file=sys.stderr)


if __name__ == "__main__":
    main()
